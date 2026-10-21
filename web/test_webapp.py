# -*- coding: utf-8 -*-
"""网页版自测：使用 Flask test client 覆盖核心接口与规则。"""
import io
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import app as webapp


def make_docx_bytes():
    """生成一个简单的 .docx（用 python-docx 构造后存字节）。"""
    import docx
    d = docx.Document()
    d.add_paragraph("深夜大雨中，落魄退伍兵陈默在便利店撞见被追杀的少女林晚。")
    d.add_paragraph("他搬空货架挡住追兵，带着少女冲进暴雨。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class WebAppTest(unittest.TestCase):

    def setUp(self):
        webapp.app.config["TESTING"] = True
        self.client = webapp.app.test_client()

    def test_index_ok(self):
        r = self.client.get("/")
        self.assertEqual(r.status_code, 200)
        html = r.get_data(as_text=True)
        self.assertIn("纯文转剧本分镜", html)
        # 三语切换器存在
        self.assertIn('data-lang="zh-CN"', html)
        self.assertIn('data-lang="zh-TW"', html)
        self.assertIn('data-lang="en"', html)

    def test_convert_text_short_compact(self):
        text = "外卖小哥获得读心术，送餐途中一次次救下陌生人。"
        r = self.client.post("/convert", data={"text": text})
        self.assertEqual(r.status_code, 200)
        data = r.get_json()
        self.assertTrue(data["ok"])
        self.assertFalse(data["break_by_shot"])
        content = data["content"]
        # 纯文字档：无表格装饰符号，台词用 角色："内容" 而非 【角色】
        self.assertNotIn("【", content)
        self.assertNotIn("|", content)
        self.assertNotIn("*", content)
        self.assertIn('"', content)
        # content 为纯分镜稿，AI 声明由前端按语言附加（后端提供 ai_notice 字段）
        self.assertNotIn("本内容由 AI 生成，仅供参考", content)
        self.assertEqual(data["ai_notice"], "本内容由 AI 生成，仅供参考")
        # 返回输入类型代码，供前端翻译标签
        self.assertIn("input_type", data)
        self.assertIn(data["input_type"], {"one_line", "summary", "excerpt", "outline"})

    def test_convert_text_long_break_by_shot(self):
        text = "深夜大雨中，落魄退伍兵陈默在便利店撞见被追杀的少女林晚。" * 40  # >800
        r = self.client.post("/convert", data={"text": text})
        self.assertEqual(r.status_code, 200)
        data = r.get_json()
        self.assertTrue(data["ok"])
        self.assertTrue(data["break_by_shot"])
        # 空行分段数量 = 分镜段数（每镜一段）
        content = data["content"]
        self.assertGreater(content.count("\n\n"), 3)

    def test_convert_upload_txt(self):
        raw = "明天太阳照常升起，小镇青年阿远决定离开家乡。"
        r = self.client.post(
            "/convert",
            data={"file": (io.BytesIO(raw.encode("utf-8")), "story.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["ok"])

    def test_convert_upload_md(self):
        raw = "# 第一章\n\n少年捡到一枚古镜，镜中映出三日后自己被杀的画面。"
        r = self.client.post(
            "/convert",
            data={"file": (io.BytesIO(raw.encode("utf-8")), "chapter.md")},
            content_type="multipart/form-data",
        )
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["ok"])

    def test_convert_upload_docx(self):
        r = self.client.post(
            "/convert",
            data={"file": (io.BytesIO(make_docx_bytes()), "story.docx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["ok"])
        self.assertIn("陈默", r.get_json()["content"])

    def test_convert_empty_input(self):
        r = self.client.post("/convert", data={"text": "   "})
        self.assertEqual(r.status_code, 400)
        data = r.get_json()
        self.assertFalse(data["ok"])
        self.assertEqual(data["error_code"], "empty_text")

    def test_convert_no_input(self):
        r = self.client.post("/convert", data={})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.get_json()["error_code"], "no_input")

    def test_convert_unsupported_ext(self):
        r = self.client.post(
            "/convert",
            data={"file": (io.BytesIO(b"x"), "evil.exe")},
            content_type="multipart/form-data",
        )
        self.assertEqual(r.status_code, 400)
        data = r.get_json()
        self.assertEqual(data["error_code"], "bad_ext")
        self.assertIn("不支持的文件类型", data["error"])
        self.assertEqual(data["detail"]["ext"], ".exe")

    def test_convert_gbk_txt(self):
        # GBK 编码的 .txt 也应能正常读取
        raw = "山城旧巷，面馆老板在深夜收留了一个说不清自己来历的女孩。"
        r = self.client.post(
            "/convert",
            data={"file": (io.BytesIO(raw.encode("gb18030")), "gbk.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(r.status_code, 200)
        data = r.get_json()
        self.assertTrue(data["ok"])
        self.assertIn("面馆老板", data["content"])

    def test_convert_oversize_file(self):
        big = b"A" * (11 * 1024 * 1024)
        r = self.client.post(
            "/convert",
            data={"file": (io.BytesIO(big), "big.txt")},
            content_type="multipart/form-data",
        )
        self.assertEqual(r.status_code, 413)
        self.assertEqual(r.get_json()["error_code"], "file_too_large")


if __name__ == "__main__":
    unittest.main(verbosity=2)